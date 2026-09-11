# -*- coding: utf-8 -*-
"""
Загрузка каталога из Excel.

    python3 tools/catalog.py template          создать import/shablon.xlsx
    python3 tools/catalog.py export            выгрузить текущий каталог в Excel
    python3 tools/catalog.py import <файл>     загрузить товары из Excel или CSV
    python3 tools/catalog.py import <файл> --build   загрузить и сразу пересобрать сайт

Файл проверяется целиком до записи: если есть хотя бы одна ошибка,
data/products.json не трогается. Флаг --force записывает только корректные строки.
"""
import csv, json, os, re, sys, shutil, datetime, subprocess

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
IMPORT_DIR = os.path.join(ROOT, "import")

GENDERS = {"men": "мужское", "women": "женское", "unisex": "унисекс"}
CATEGORIES = {"running": "бег", "lifestyle": "lifestyle", "training": "тренировки"}
COLORS = ["black", "white", "grey", "red", "green", "pink", "brown", "blue",
          "beige", "yellow", "orange", "purple", "multi"]

# (ключ в json, заголовок в Excel, обязательное, ширина колонки)
COLUMNS = [
    ("id",          "ID",                 False, 30),
    ("brand",       "Бренд",              True,  16),
    ("model",       "Модель",             True,  26),
    ("sku",         "Артикул",            True,  16),
    ("styleCode",   "Код производителя",  False, 18),
    ("gender",      "Пол",                True,  12),
    ("category",    "Категория",          True,  14),
    ("purpose",     "Назначение",         True,  30),
    ("color",       "Цвет (код)",         True,  13),
    ("colorName",   "Цвет (название)",    True,  24),
    ("price",       "Цена",               True,  11),
    ("oldPrice",    "Старая цена",        False, 13),
    ("isNew",       "Новинка",            False, 10),
    ("releasedAt",  "Дата поступления",   False, 18),
    ("upper",       "Материал верха",     True,  26),
    ("sole",        "Материал подошвы",   True,  26),
    ("country",     "Страна",             True,  14),
    ("image",       "Изображение",        True,  34),
    ("description", "Описание",           True,  60),
    ("sizes",       "Размеры",            True,  34),
]

HEADER_BY_KEY = {k: h for k, h, _, _ in COLUMNS}
KEY_BY_HEADER = {h.lower(): k for k, h, _, _ in COLUMNS}
REQUIRED = [k for k, _, req, _ in COLUMNS if req]


def die(msg):
    print("Ошибка: " + msg)
    sys.exit(1)


def load(name):
    return json.load(open(os.path.join(DATA, name), encoding="utf-8"))


# ── чтение ────────────────────────────────────────────────────────────

def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt"):
        return _read_csv(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    die("поддерживаются .xlsx и .csv, получен «%s»" % (ext or "файл без расширения"))


def _headers_to_keys(header_row):
    keys, unknown = [], []
    for cell in header_row:
        # звёздочка отмечает обязательные колонки — при сопоставлении её отбрасываем
        h = re.sub(r"[\s*]+$", "", str(cell or "").strip()).lower()
        if h in KEY_BY_HEADER:
            keys.append(KEY_BY_HEADER[h])
        elif h in HEADER_BY_KEY:          # допускаем англоязычные ключи
            keys.append(h)
        elif h:
            keys.append(None)
            unknown.append(str(cell).strip())
        else:
            keys.append(None)
    return keys, unknown


def _rows_from(table, keys):
    out = []
    for i, row in enumerate(table, start=2):
        d = {"__row": i}
        for k, v in zip(keys, row):
            if k:
                d[k] = v
        if any(str(v).strip() for k, v in d.items() if k != "__row" and v is not None):
            out.append(d)
    return out


def _read_csv(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(f, dialect))
    if not rows:
        die("файл пуст")
    keys, unknown = _headers_to_keys(rows[0])
    if unknown:
        print("Колонки не распознаны и пропущены: " + ", ".join(unknown))
    return _rows_from(rows[1:], keys)


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        die("для .xlsx нужен openpyxl: pip3 install openpyxl\n"
            "       (или сохраните файл как CSV — он читается без библиотек)")
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Товары"] if "Товары" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        die("лист пуст")
    keys, unknown = _headers_to_keys(rows[0])
    if unknown:
        print("Колонки не распознаны и пропущены: " + ", ".join(unknown))
    return _rows_from(rows[1:], keys)


# ── разбор значений ───────────────────────────────────────────────────

def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def parse_bool(v):
    return s(v).lower() in ("да", "yes", "true", "1", "+", "новинка")


def parse_money(v, field, errs):
    t = s(v).replace(" ", "").replace(" ", "").replace("₽", "").replace(",", ".")
    if not t:
        return None
    try:
        n = float(t)
    except ValueError:
        errs.append("%s: «%s» — не число" % (field, s(v)))
        return None
    if n <= 0:
        errs.append("%s должна быть больше нуля" % field)
        return None
    return int(round(n))


def _fmt_size(x):
    return str(int(x)) if float(x) == int(x) else str(x)


def parse_sizes(v, errs):
    """«40:3, 41:4, 42:0» → [{"size":40,"stock":3}, …]. Без остатка считаем 0."""
    raw = s(v)
    if not raw:
        errs.append("Размеры: пусто")
        return []
    out, seen = [], set()
    for part in re.split(r"[,;\n]+", raw):
        part = part.strip()
        if not part:
            continue
        m = re.match(r"^(\d+(?:[.,]\d+)?)\s*(?::\s*(\d+))?$", part)
        if not m:
            errs.append("Размеры: «%s» — ожидается «размер:остаток», например 42:3" % part)
            continue
        size = float(m.group(1).replace(",", "."))
        stock = int(m.group(2)) if m.group(2) else 0
        if size in seen:
            errs.append("Размеры: %s повторяется" % _fmt_size(size))
            continue
        seen.add(size)
        out.append({"size": int(size) if size == int(size) else size, "stock": stock})
    out.sort(key=lambda x: float(x["size"]))
    return out


def slugify(*parts):
    tr = {"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
          "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
          "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c",
          "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e",
          "ю": "yu", "я": "ya"}
    t = " ".join(p for p in parts if p).lower()
    t = "".join(tr.get(ch, ch) for ch in t)
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return re.sub(r"-{2,}", "-", t)


# ── проверка ──────────────────────────────────────────────────────────

def validate(rows, brands):
    products, problems = [], []
    seen_id, seen_sku = {}, {}
    brand_slugs = {b["slug"]: b for b in brands}
    brand_by_name = {b["name"].lower(): b["slug"] for b in brands}

    for r in rows:
        line = r.get("__row", "?")
        errs, warns = [], []

        for key in REQUIRED:
            if not s(r.get(key)):
                errs.append("%s: не заполнено" % HEADER_BY_KEY[key])

        brand = s(r.get("brand")).lower()
        if brand and brand not in brand_slugs:
            if brand in brand_by_name:
                brand = brand_by_name[brand]
            else:
                errs.append("Бренд: «%s» нет в data/brands.json — добавьте бренд или исправьте название"
                            % s(r.get("brand")))

        gender = s(r.get("gender")).lower()
        if gender in GENDERS.values():
            gender = [k for k, v in GENDERS.items() if v == gender][0]
        if gender and gender not in GENDERS:
            errs.append("Пол: «%s» — допустимо %s" % (s(r.get("gender")), ", ".join(GENDERS)))

        category = s(r.get("category")).lower()
        if category in CATEGORIES.values():
            category = [k for k, v in CATEGORIES.items() if v == category][0]
        if category and category not in CATEGORIES:
            errs.append("Категория: «%s» — допустимо %s" % (s(r.get("category")), ", ".join(CATEGORIES)))

        color = s(r.get("color")).lower()
        if color and color not in COLORS:
            warns.append("Цвет «%s» не из списка — фильтр по цвету его не покажет" % color)

        price = parse_money(r.get("price"), "Цена", errs)
        old = parse_money(r.get("oldPrice"), "Старая цена", errs) if s(r.get("oldPrice")) else None
        if price and old and old <= price:
            errs.append("Старая цена (%s) должна быть больше текущей (%s)" % (old, price))

        sizes = parse_sizes(r.get("sizes"), errs)

        raw_date = r.get("releasedAt")
        if isinstance(raw_date, datetime.datetime):
            released = raw_date.date().isoformat()
        elif isinstance(raw_date, datetime.date):
            released = raw_date.isoformat()
        else:
            released = s(raw_date) or datetime.date.today().isoformat()
            if not re.match(r"^\d{4}-\d{2}-\d{2}$", released):
                errs.append("Дата поступления: «%s» — ожидается ГГГГ-ММ-ДД" % released)

        pid = s(r.get("id")) or slugify(brand, s(r.get("model")), s(r.get("colorName")))
        sku = s(r.get("sku"))
        if pid in seen_id:
            errs.append("ID «%s» уже использован в строке %s" % (pid, seen_id[pid]))
        if sku and sku in seen_sku:
            errs.append("Артикул «%s» уже использован в строке %s" % (sku, seen_sku[sku]))

        if errs:
            problems.append((line, s(r.get("model")) or "—", errs, warns))
            continue

        seen_id[pid] = line
        if sku:
            seen_sku[sku] = line
        if warns:
            problems.append((line, s(r.get("model")), [], warns))

        products.append({
            "id": pid, "brand": brand, "model": s(r.get("model")),
            "sku": sku, "styleCode": s(r.get("styleCode")),
            "gender": gender, "category": category, "purpose": s(r.get("purpose")),
            "color": color, "colorName": s(r.get("colorName")),
            "price": price, "oldPrice": old,
            "isNew": parse_bool(r.get("isNew")), "releasedAt": released,
            "materials": {"upper": s(r.get("upper")), "sole": s(r.get("sole"))},
            "country": s(r.get("country")),
            "image": s(r.get("image")),
            "description": s(r.get("description")),
            "sizes": sizes,
        })
    return products, problems


# ── запись Excel ──────────────────────────────────────────────────────

def _to_row(p):
    return {
        "id": p["id"], "brand": p["brand"], "model": p["model"], "sku": p["sku"],
        "styleCode": p.get("styleCode", ""), "gender": p["gender"],
        "category": p["category"], "purpose": p["purpose"], "color": p["color"],
        "colorName": p["colorName"], "price": p["price"],
        "oldPrice": p.get("oldPrice") or "", "isNew": "да" if p.get("isNew") else "",
        "releasedAt": p["releasedAt"], "upper": p["materials"]["upper"],
        "sole": p["materials"]["sole"], "country": p["country"],
        "image": p["image"], "description": p["description"],
        "sizes": ", ".join("%s:%d" % (_fmt_size(x["size"]), x["stock"]) for x in p["sizes"]),
    }


def write_xlsx(path, products, brands):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError:
        die("для записи .xlsx нужен openpyxl: pip3 install openpyxl")

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"

    head_font = Font(bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor="111111")
    req_fill = PatternFill("solid", fgColor="3B1F2B")

    for i, (key, header, req, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=header + (" *" if req else ""))
        c.font = head_font
        c.fill = req_fill if req else head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    for r, p in enumerate(products, start=2):
        vals = _to_row(p)
        for i, (key, _, _, _) in enumerate(COLUMNS, start=1):
            ws.cell(row=r, column=i, value=vals.get(key, ""))

    last = max(len(products) + 1, 200)
    order = [k for k, _, _, _ in COLUMNS]

    def col(key):
        return get_column_letter(order.index(key) + 1)

    for key, allowed in (("gender", list(GENDERS)), ("category", list(CATEGORIES)),
                         ("color", COLORS), ("isNew", ["да"])):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(allowed), allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("%s2:%s%d" % (col(key), col(key), last))

    dvb = DataValidation(type="list",
                         formula1='"%s"' % ",".join(b["slug"] for b in brands), allow_blank=True)
    ws.add_data_validation(dvb)
    dvb.add("%s2:%s%d" % (col("brand"), col("brand"), last))

    hs = wb.create_sheet("Справка")
    hs.column_dimensions["A"].width = 24
    hs.column_dimensions["B"].width = 96
    rows = [
        ("Как пользоваться", ""),
        ("", "Заполните лист «Товары» и выполните: python3 tools/catalog.py import <файл> --build"),
        ("", "Колонки со звёздочкой обязательны. Порядок колонок можно менять, лишние игнорируются."),
        ("", "Импорт полностью заменяет каталог: что есть в файле, то и будет на сайте."),
        ("", ""),
        ("ID", "Можно не заполнять — соберётся из бренда, модели и цвета. Участвует в адресе страницы."),
        ("Бренд", "Код из data/brands.json: " + ", ".join(b["slug"] for b in brands)),
        ("Артикул", "Ваш внутренний номер, уникальный. По нему работает поиск."),
        ("Код производителя", "Необязательно. Артикул бренда, если нужен."),
        ("Пол", "men, women или unisex. Унисекс попадает и в мужской, и в женский раздел."),
        ("Категория", "running, lifestyle или training."),
        ("Назначение", "Свободный текст, например «Бег по асфальту». Работает как фильтр."),
        ("Цвет (код)", "Один из: " + ", ".join(COLORS) + ". По нему работает фильтр цвета."),
        ("Цвет (название)", "Как показывать покупателю, например «Black / Gym Red»."),
        ("Цена", "Число в рублях, без пробелов и знака ₽."),
        ("Старая цена", "Необязательно. Если больше текущей — появится скидка и товар попадёт в Sale."),
        ("Новинка", "«да» — товар попадёт в раздел «Новинки» и получит метку."),
        ("Дата поступления", "ГГГГ-ММ-ДД. Используется для сортировки «Новинки»."),
        ("Изображение", "ID фотографии Unsplash (photo-…) либо путь к своему файлу, "
                        "например assets/img/model.jpg"),
        ("Размеры", "«размер:остаток» через запятую: 40:3, 41:4, 42:0. "
                    "Остаток 0 — размер показывается зачёркнутым."),
        ("", ""),
        ("Проверка", "Если в файле есть ошибки, каталог не перезаписывается — исправьте и запустите снова."),
        ("Резервная копия", "Перед записью сохраняется data/products.backup.json."),
    ]
    for i, (a, b_) in enumerate(rows, start=1):
        ca = hs.cell(row=i, column=1, value=a)
        ca.font = Font(bold=True)
        hs.cell(row=i, column=2, value=b_).alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def write_csv(path, products):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h + (" *" if req else "") for _, h, req, _ in COLUMNS])
        for p in products:
            v = _to_row(p)
            w.writerow([v.get(k, "") for k, _, _, _ in COLUMNS])


# ── команды ───────────────────────────────────────────────────────────

def cmd_template(fmt):
    brands = load("brands.json")
    sample = [{
        "id": "", "brand": brands[0]["slug"], "model": "Gel-Kayano 31",
        "sku": "FRM-XX-0001", "styleCode": "", "gender": "men", "category": "running",
        "purpose": "Бег по асфальту", "color": "blue", "colorName": "Blue / Silver",
        "price": 16990, "oldPrice": None, "isNew": True,
        "releasedAt": datetime.date.today().isoformat(),
        "materials": {"upper": "Инженерная сетка", "sole": "Резина, вставка из пены"},
        "country": "Вьетнам", "image": "photo-0000000000000-000000000000",
        "description": "Пример строки. Удалите её и впишите свои товары.",
        "sizes": [{"size": 41, "stock": 2}, {"size": 42, "stock": 3}, {"size": 43, "stock": 0}],
    }]
    if fmt == "csv":
        path = os.path.join(IMPORT_DIR, "shablon.csv")
        write_csv(path, sample)
    else:
        path = os.path.join(IMPORT_DIR, "shablon.xlsx")
        write_xlsx(path, sample, brands)
    print("Шаблон: %s" % os.path.relpath(path, ROOT))
    print("Заполните лист «Товары», на листе «Справка» — что значит каждая колонка.")


def cmd_export(fmt):
    products, brands = load("products.json"), load("brands.json")
    if fmt == "csv":
        path = os.path.join(IMPORT_DIR, "products.csv")
        write_csv(path, products)
    else:
        path = os.path.join(IMPORT_DIR, "products.xlsx")
        write_xlsx(path, products, brands)
    rel = os.path.relpath(path, ROOT)
    print("Выгружено %d товаров → %s" % (len(products), rel))
    print("Правьте файл и загружайте обратно: python3 tools/catalog.py import %s --build" % rel)


def cmd_import(path, force=False, build=False):
    if not os.path.isfile(path):
        die("файл не найден: %s" % path)
    brands = load("brands.json")
    rows = read_rows(path)
    if not rows:
        die("в файле нет строк с товарами")
    products, problems = validate(rows, brands)

    errors = [p for p in problems if p[2]]
    warns = [p for p in problems if p[3]]

    if warns:
        print("\nПредупреждения:")
        for line, model, _, ws in warns:
            for w in ws:
                print("  строка %s (%s): %s" % (line, model, w))

    if errors:
        print("\nОшибки — %d строк(и):" % len(errors))
        for line, model, es, _ in errors:
            print("  строка %s (%s):" % (line, model or "без модели"))
            for e_ in es:
                print("      • " + e_)
        if not force:
            print("\nКаталог не изменён. Исправьте файл и запустите снова "
                  "или добавьте --force, чтобы загрузить только корректные строки.")
            sys.exit(1)
        print("\n--force: пропускаем строки с ошибками.")

    if not products:
        die("нет ни одной корректной строки")

    target = os.path.join(DATA, "products.json")
    if os.path.exists(target):
        shutil.copy(target, os.path.join(DATA, "products.backup.json"))
    with open(target, "w", encoding="utf-8") as f:
        json.dump(products, f, ensure_ascii=False, indent=2)
        f.write("\n")

    in_stock = sum(1 for p in products if sum(x["stock"] for x in p["sizes"]) > 0)
    print("\nЗагружено товаров: %d (в наличии: %d, под заказ: %d)"
          % (len(products), in_stock, len(products) - in_stock))
    print("Резервная копия прежнего каталога: data/products.backup.json")

    if build:
        print("\nПересборка сайта…")
        subprocess.run([sys.executable, os.path.join(ROOT, "build.py")], cwd=ROOT, check=True)
    else:
        print("Теперь выполните: python3 build.py")


def main():
    args = sys.argv[1:]
    if not args or args[0] in ("-h", "--help", "help"):
        print(__doc__)
        return
    cmd, fmt = args[0], ("csv" if "--csv" in args else "xlsx")
    if cmd == "template":
        cmd_template(fmt)
    elif cmd == "export":
        cmd_export(fmt)
    elif cmd == "import":
        files = [a for a in args[1:] if not a.startswith("-")]
        if not files:
            die("укажите файл: python3 tools/catalog.py import import/products.xlsx")
        cmd_import(files[0], force="--force" in args, build="--build" in args)
    else:
        die("неизвестная команда «%s». Доступны: template, export, import" % cmd)


if __name__ == "__main__":
    main()
